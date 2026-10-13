#!/usr/bin/env python3
"""Match OpenStreetMap church buildings to registered church charities by POSTCODE (not by name).

For each OSM church with a postcode, find charities on the team list with the same full postcode:
  postcode_unique     exactly one charity at that postcode
  postcode+name       several charities there; one shares a distinctive word with the OSM name
  postcode_ambiguous  several charities there and the name does not pick one (all candidates listed)
  no_charity_at_postcode / no_postcode
Name and website are used only as supporting evidence (name_also_matches, website_also_matches).
Caveat: the charity's postcode is its CONTACT address (sometimes a treasurer's home), not always the building.

Usage: .venv/bin/python scripts/match_osm_postcodes.py --sheet "<Bankwell Data Analysis.xlsx>" --out <csv>
"""

from __future__ import annotations

import argparse
import collections
import csv
import re
import sys

COMMON = {"church", "chapel", "saint", "st", "the", "of", "and", "parish", "united", "free", "hall", "centre",
          "christian", "fellowship", "ministries", "mission", "trust", "charity", "council", "parochial",
          "ecclesiastical", "baptist", "methodist", "catholic", "roman", "anglican", "evangelical", "pentecostal",
          "reformed", "congregational", "community", "new", "life", "all", "saints", "mary", "virgin", "christ"}


def norm_pc(pc) -> str:
    return re.sub(r"\s+", "", str(pc or "")).upper()


def words(name) -> set[str]:
    return {w for w in re.findall(r"[a-z]{3,}", str(name or "").lower()) if w not in COMMON}


def domain(url) -> str:
    m = re.match(r"^(?:https?://)?(?:www\.)?([^/\s?#]+)", str(url or "").strip().lower())
    return m.group(1) if m and "." in m.group(1) else ""


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--sheet", required=True)
    p.add_argument("--out", required=True)
    a = p.parse_args()
    import openpyxl
    wb = openpyxl.load_workbook(a.sheet, read_only=True, data_only=True)
    main_rows = wb["UK Churches (from python script"].iter_rows(values_only=True)
    mh = next(main_rows)
    mi = {k: i for i, k in enumerate(mh)}
    by_pc = collections.defaultdict(list)
    for r in main_rows:
        pc = norm_pc(r[mi["contact_postcode"]])
        if pc:
            by_pc[pc].append(r)
    osm_rows = wb["All UK Churches (OpenStreetMap)"].iter_rows(values_only=True)
    oh = next(osm_rows)
    oi = {k: i for i, k in enumerate(oh)}

    stats = collections.Counter()
    out_fields = list(oh) + ["match_type", "charity_number", "charity_name", "charity_postcode", "candidates_at_postcode",
                             "name_also_matches", "website_also_matches", "all_candidates"]
    with open(a.out, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(out_fields)
        for r in osm_rows:
            pc = norm_pc(r[oi["addr:postcode"]])
            site = domain(r[oi["website"]] or r[oi["contact:website"]])
            cands = by_pc.get(pc, []) if pc else []
            chosen, kind = None, ""
            if not pc:
                kind = "no_postcode"
            elif not cands:
                kind = "no_charity_at_postcode"
            elif len(cands) == 1:
                chosen, kind = cands[0], "postcode_unique"
            else:
                ow = words(r[oi["name"]])
                hits = [c for c in cands if ow & words(c[mi["charity_name"]])]
                if len(hits) == 1:
                    chosen, kind = hits[0], "postcode+name"
                else:
                    kind = "postcode_ambiguous"
            stats[kind] += 1
            name_ok = bool(chosen) and bool(words(r[oi["name"]]) & words(chosen[mi["charity_name"]]))
            web_ok = bool(chosen) and bool(site) and site == domain(chosen[mi["contact_web"]])
            w.writerow(list(r) + [
                kind,
                chosen[mi["registered_number"]] if chosen else "",
                chosen[mi["charity_name"]] if chosen else "",
                chosen[mi["contact_postcode"]] if chosen else "",
                len(cands),
                "yes" if name_ok else ("no" if chosen else ""),
                "yes" if web_ok else ("no" if chosen else ""),
                " | ".join(f"{c[mi['registered_number']]}: {c[mi['charity_name']]}" for c in cands) if kind == "postcode_ambiguous" else "",
            ])
            if chosen:
                stats["matched_name_also"] += name_ok
                stats["matched_web_also"] += web_ok
    print(f"saved {a.out}")
    for k, v in stats.most_common():
        print(f"  {k}: {v}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
