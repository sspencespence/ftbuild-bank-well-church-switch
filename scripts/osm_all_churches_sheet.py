#!/usr/bin/env python3
"""All UK church buildings from OpenStreetMap, with the sort code / bank found on each one's own website.

Input: $BANKWELL_DATA/osm-crawl/osm_all.csv (Overpass, scripts/osm/q_all.overpassql), the OSM website crawl
(osm-crawl/church_rows.json + sort_codes.jsonl) and sortcode-banks/sort_code_banks.csv (see find-church-banks).
`kind`: "place of worship" (amenity=place_of_worship + religion=christian) or "<building> building (not tagged
as worship)" - may be disused or converted. Names/postcodes/websites are as entered by OSM volunteers.

Usage: .venv/bin/python scripts/osm_all_churches_sheet.py [all_uk_churches_osm.xlsx] [results/osm_all_uk_churches.csv]
"""
import collections
import csv
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent))
from bankwell_paths import DATA  # noqa: E402

O = DATA / "osm-crawl"
HDR = ["osm_type", "osm_id", "name", "denomination", "kind", "website", "postcode", "city", "lat", "lon", "sort_code",
       "bank (from sort code)"]
README = ["All UK churches from OpenStreetMap — pulled 10 Oct 2026",
          "Includes: every amenity=place_of_worship + religion=christian in the UK, plus buildings tagged church/chapel/cathedral "
          "that are not tagged as a place of worship (may be disused or converted).",
          "Names, denominations, postcodes and websites are as entered by OpenStreetMap volunteers — often missing. No charity numbers "
          "(see osm_churches_linked for the link to the charity register).",
          "Sort code / bank (yellow) come from crawling the church website, where OSM has one (see osm_church_sort_codes.xlsx for evidence).",
          "Map data © OpenStreetMap contributors (ODbL)."]


def build() -> tuple[list[list], collections.Counter]:
    codes = {r["sort_code"]: r["bank"] for r in csv.DictReader(open(DATA / "sortcode-banks/sort_code_banks.csv"))}
    recs = {}
    for line in open(O / "sort_codes.jsonl"):
        d = json.loads(line)
        recs[d["regno"]] = d
    sc = {}
    for r in json.load(open(O / "church_rows.json")):
        d = recs.get(r["site_key"], {})
        if d.get("sort_codes"):
            c = d["sort_codes"][0]
            sc[(r["osm_type"], str(r["osm_id"]))] = (c, codes.get(c, "unknown") or "unknown")
    out, st = [], collections.Counter()
    for r in csv.DictReader(open(O / "osm_all.csv", encoding="utf-8")):
        kind = "place of worship" if r["amenity"] == "place_of_worship" else f"{r['building']} building (not tagged as worship)"
        web = r["website"] or r["contact:website"] or r["url"]
        s = sc.get((r["@type"], r["@id"]), ("", ""))
        out.append([r["@type"], int(r["@id"]), r["name"], r["denomination"], kind, web, r["addr:postcode"], r["addr:city"],
                    float(r["@lat"]) if r["@lat"] else None, float(r["@lon"]) if r["@lon"] else None, s[0], s[1]])
        st[kind.split(" (")[0]] += 1
        st["with website"] += bool(web)
        st["with postcode"] += bool(r["addr:postcode"])
        st["with sort code"] += bool(s[0])
    return out, st


def main() -> int:
    xlsx = sys.argv[1] if len(sys.argv) > 1 else "all_uk_churches_osm.xlsx"
    out, st = build()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All UK churches (OSM)"
    ws.append(HDR)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="325569")
    y = PatternFill("solid", fgColor="FFF2A8")
    for row in out:
        ws.append(row)
        if row[10]:
            ws.cell(ws.max_row, 11).fill = y
            ws.cell(ws.max_row, 12).fill = y
    for col, w in zip("ABCDEFGHIJKL", [9, 13, 34, 16, 30, 36, 10, 14, 10, 10, 11, 20]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    n = wb.create_sheet("READ ME", 0)
    for t in README:
        n.append([t])
    n.column_dimensions["A"].width = 140
    n["A1"].font = Font(bold=True, size=13)
    wb.save(xlsx)
    if len(sys.argv) > 2:
        with open(sys.argv[2], "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(HDR)
            w.writerows(out)
    print(len(out), dict(st))
    return 0


if __name__ == "__main__":
    sys.exit(main())
