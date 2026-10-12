#!/usr/bin/env python3
"""Write the PDF-extraction result (.csv or .xlsx): registered_number, assets_cash, figures_source.

Same acceptance rules as merge_into_sheet.py (a figure is only written when it passed the checks).
One row per church whose PDF has been read, in spreadsheet order. Churches with official figures
are not included (they are already in the team sheet).

Usage:
    .venv/bin/python scripts/export_pdf_csv.py --sheet "Bankwell Data Analysis.xlsx" \
        --out results/pdf_extract_final.csv
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from merge_into_sheet import E, FIELDS, TAB, decide, load_rows  # noqa: E402


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--sheet", required=True)
    p.add_argument("--out", required=True)
    a = p.parse_args()
    import openpyxl
    first = load_rows(str(E / "out" / "target-*.jsonl"))
    second = load_rows(str(E / "check" / "*.jsonl"))
    docs = {d["regno"] for d in map(json.loads, (E / "docs.jsonl").read_text().splitlines())}
    rows = openpyxl.load_workbook(a.sheet, read_only=True, data_only=True)[TAB].iter_rows(values_only=True)
    head = next(rows)
    i_reg, i_cash = head.index("registered_number"), head.index("assets_cash")
    lines, n_cash = ["registered_number\tassets_cash\tfigures_source"], 0
    for r in rows:
        regno = str(r[i_reg] or "").split(".")[0]
        if r[i_cash] not in (None, "") or regno not in docs:
            continue
        if not any((regno, f) in first for f in FIELDS):
            continue  # not read yet
        values = {f: decide(first.get((regno, f)), second.get((regno, f)))[0] for f in FIELDS}
        cash = values["assets_cash"]
        source = "PDF accounts read by Claude" if any(v is not None for v in values.values()) else "PDF read - figures not found"
        shown = "" if cash is None else (f"{cash:.2f}".rstrip("0").rstrip(".") if cash % 1 else str(int(cash)))
        n_cash += cash is not None
        lines.append(f"{regno}\t{shown}\t{source}")
    if a.out.lower().endswith(".xlsx"):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "pdf_extract"
        yellow = PatternFill("solid", fgColor="FFFF00")
        ws.append(lines[0].split("\t"))
        for line in lines[1:]:
            reg, cash, source = line.split("\t")
            ws.append([int(reg) if reg.isdigit() else reg, float(cash) if cash else None, source])
            if cash:
                ws.cell(ws.max_row, 2).fill = yellow
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 18, 14, 32
        ws.freeze_panes = "A2"
        wb.save(a.out)
    else:
        import csv
        with open(a.out, "w", encoding="utf-8", newline="") as fh:
            csv.writer(fh).writerows(line.split("\t") for line in lines)
    print(f"wrote {a.out}: {len(lines) - 1} churches read, {n_cash} with cash")
    return 0


if __name__ == "__main__":
    sys.exit(main())
